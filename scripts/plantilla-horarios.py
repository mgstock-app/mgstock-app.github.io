"""Genera las planillas de Horarios de GESTLOK (mismo formato semanal que usa el local de Munro).

  public/plantillas/GESTLOK-plantilla-horarios.xlsx       plantilla predeterminada, para completar
  public/plantillas/GESTLOK-demo-horarios-ficticio.xlsx   ejemplo completo con datos ficticios (demo)

Uso: python3 scripts/plantilla-horarios.py   (después: recalcular con LibreOffice para guardar valores)
"""
import json
from datetime import date, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public" / "plantillas"
DEMO = json.loads((OUT / "demo-horarios.json").read_text(encoding="utf-8"))

F = "Arial"
NAVY, INK, MUTED = "123462", "17263E", "5D6B7E"
INPUT = PatternFill("solid", fgColor="E3F2FD")      # celeste: completar
HEAD = PatternFill("solid", fgColor=NAVY)
SOFT = PatternFill("solid", fgColor="F3F6FA")
TOTAL = PatternFill("solid", fgColor="EAF6EE")
thin = Side(style="thin", color="C9D3DF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
ESTADOS = ["OFF", "Feriado", "Ausente", "Ausente justificado", "Vacaciones"]
PEOPLE = 10
DAYS = "CDEFGHI"


def font(**kw):
    return Font(name=F, color=kw.pop("color", INK), **kw)


def hm(text):
    if not text:
        return None
    h, m = map(int, text.split(":"))
    return time(h, m)


def week_sheet(ws, local, monday, team, week):
    ws.title = "Semana"
    ws.sheet_view.showGridLines = False
    widths = {"A": 20, "B": 17, "J": 11, "K": 12, "L": 8, "M": 11, "N": 9, "O": 17, "P": 3}
    for c in DAYS:
        widths[c] = 11
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    for c in "QRSTUV":
        ws.column_dimensions[c].width = 11
        ws.column_dimensions[c].outlineLevel = 1
        ws.column_dimensions[c].hidden = True
    ws.merge_cells("A1:O1")
    ws["A1"] = "Planilla semanal de horarios · GESTLOK"
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    ws["D2"], ws["G2"], ws["J2"], ws["M2"] = "Local", "Semana desde (lunes)", "hasta", "Celeste = completar"
    for ref in ("D2", "G2", "J2"):
        ws[ref].font = font(bold=True, size=10, color=MUTED)
    ws["M2"].font = font(italic=True, size=9, color=MUTED)
    ws["E2"] = local
    ws["I2"] = monday
    ws["I2"].number_format = "dd/mm/yyyy"
    ws["K2"] = "=I2+6"
    ws["K2"].number_format = "dd/mm/yyyy"
    for ref in ("E2", "I2"):
        ws[ref].fill = INPUT
        ws[ref].font = font(bold=True)
        ws[ref].border = BOX
    ws["K2"].font = font(bold=True)

    heads = {"A4": "Colaborador", "B4": "Detalle", "J4": "Horas\ntrabajadas", "K4": "Turnos", "L4": "OFF\n(días)",
             "M4": "Feriado", "N4": "Ausentes\n(días)", "O4": "Total a liquidar\n(trab. + OFF + fer.)",
             "Q4": "Cobra por", "R4": "Horas OFF", "S4": "Horas en feriado", "T4": "Turnos en feriado",
             "U4": "Total horas", "V4": "Total turnos"}
    for ref, txt in heads.items():
        ws[ref] = txt
    for i, c in enumerate(DAYS):
        ws[f"{c}4"] = f'=CHOOSE(WEEKDAY({c}5,2),"Lunes","Martes","Miércoles","Jueves","Viernes","Sábado","Domingo")'
        ws[f"{c}5"] = "=$I$2" if i == 0 else f"={DAYS[i - 1]}5+1"
        ws[f"{c}5"].number_format = "dd/mm"
        ws[f"{c}6"] = (f'=IF(COUNTIFS(Feriados!$A$5:$A$60,{c}5,Feriados!$C$5:$C$60,"Feriado")>0,"Feriado","")')
    for row in (4, 5, 6):
        for col in "ABCDEFGHIJKLMNO":
            cell = ws[f"{col}{row}"]
            cell.fill = HEAD if row == 4 else SOFT
            cell.font = font(bold=True, size=10, color="FFFFFF" if row == 4 else NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BOX
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "C7"

    estado_dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=True,
                               showErrorMessage=True, errorTitle="Estado", error="Elegí un estado de la lista o dejalo vacío.")
    time_dv = DataValidation(type="time", operator="between", formula1="0", formula2="0.999988", allow_blank=True,
                             showErrorMessage=True, errorTitle="Horario", error="Escribí la hora con formato 10:00.")
    ws.add_data_validation(estado_dv)
    ws.add_data_validation(time_dv)
    labels = ["Mañana · entrada", "Mañana · salida", "Noche · entrada", "Noche · salida", "Estado", "Horas del día"]
    for i in range(PEOPLE):
        r = 7 + 7 * i
        n = i + 1
        ws[f"A{r}"] = f'=IFERROR(INDEX(Equipo!$B$5:$B$14,{n})&"","")'
        ws[f"A{r + 1}"] = f'=IF($A{r}="","",INDEX(Equipo!$C$5:$C$14,{n})&"")'
        ws[f"A{r + 2}"] = f'=IF($A{r}="","","Cobra por "&LOWER(INDEX(Equipo!$D$5:$D$14,{n})))'
        ws[f"A{r}"].font = font(bold=True, size=11)
        for k in (1, 2):
            ws[f"A{r + k}"].font = font(size=9, color=MUTED)
        for k, label in enumerate(labels):
            ws[f"B{r + k}"] = label
            ws[f"B{r + k}"].font = font(size=9, color=MUTED, bold=(k == 5))
        for c in DAYS:
            for k in range(4):
                cell = ws[f"{c}{r + k}"]
                cell.fill, cell.border, cell.number_format = INPUT, BOX, "hh:mm"
                cell.alignment = Alignment(horizontal="center")
                time_dv.add(cell)
            st = ws[f"{c}{r + 4}"]
            st.fill, st.border = INPUT, BOX
            st.font = font(size=9, bold=True)
            st.alignment = Alignment(horizontal="center", shrink_to_fit=True)
            estado_dv.add(st)
            h = ws[f"{c}{r + 5}"]
            h.value = (f'=IF(AND({c}{r + 4}<>"",{c}{r + 4}<>"Feriado"),0,'
                       f'IF(AND({c}{r}<>"",{c}{r + 1}<>""),MOD({c}{r + 1}-{c}{r},1),0)+'
                       f'IF(AND({c}{r + 2}<>"",{c}{r + 3}<>""),MOD({c}{r + 3}-{c}{r + 2},1),0))')
            h.number_format = "[h]:mm;;–"
            h.font = font(size=9, color=MUTED)
            h.alignment = Alignment(horizontal="center")
            h.border = BOX
        d, e, st_row, hrow = f"C{r}:I{r}", f"C{r + 2}:I{r + 2}", f"C{r + 4}:I{r + 4}", f"C{r + 5}:I{r + 5}"
        ws[f"J{r}"] = f'=IF($Q{r}="Turnos","",SUM({hrow}))'
        ws[f"K{r}"] = f'=SUMPRODUCT((({d}<>"")+({e}<>""))*(({st_row}="")+({st_row}="Feriado")))'
        ws[f"L{r}"] = f'=COUNTIF({st_row},"OFF")'
        ws[f"M{r}"] = (f'=IF($A{r}="","",IF($Q{r}="Horas",IF($S{r}=0,"–",TEXT($S{r},"[h]:mm")&" hs"),'
                       f'IF($T{r}=0,"–",$T{r}&IF($T{r}=1," turno"," turnos"))))')
        ws[f"N{r}"] = f'=COUNTIF({st_row},"Ausente")'
        ws[f"O{r}"] = f'=IF($A{r}="","",IF($Q{r}="Horas",TEXT($U{r},"[h]:mm")&" hs",$V{r}&IF($V{r}=1," turno"," turnos")))'
        ws[f"Q{r}"] = f'=IFERROR(INDEX(Equipo!$D$5:$D$14,{n})&"","")'
        ws[f"R{r}"] = f"=N(INDEX(Equipo!$E$5:$E$14,{n}))"
        ws[f"S{r}"] = f'=SUMPRODUCT({hrow}*({st_row}="Feriado"))'
        ws[f"T{r}"] = f'=COUNTIFS({st_row},"Feriado",{d},"<>")+COUNTIFS({st_row},"Feriado",{e},"<>")'
        ws[f"U{r}"] = f"=N(J{r})+L{r}*R{r}+S{r}"
        ws[f"V{r}"] = f"=K{r}+L{r}+T{r}"
        ws[f"J{r}"].number_format = "[h]:mm;;–"
        for col in "JKLMNO":
            ws.merge_cells(f"{col}{r}:{col}{r + 5}")
            cell = ws[f"{col}{r}"]
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = font(bold=(col == "O"), size=11 if col == "O" else 10)
            cell.fill = TOTAL if col == "O" else PatternFill()
            cell.border = BOX
        if i < len(team):
            for day, (m_in, m_out, n_in, n_out, estado) in enumerate(week.get(team[i]["nombre"], [])):
                c = DAYS[day]
                for k, v in enumerate((m_in, m_out, n_in, n_out)):
                    ws[f"{c}{r + k}"] = hm(v)
                ws[f"{c}{r + 4}"] = estado or None
    last = 7 + 7 * PEOPLE
    ws.merge_cells(f"A{last}:O{last}")
    ws[f"A{last}"] = ("Turnos: cada horario de entrada cargado suma 1 turno. Horas: salida − entrada de cada turno "
                      "(si sale pasada la medianoche se calcula bien). Un OFF suma las horas de la pestaña Equipo o 1 turno.")
    ws[f"A{last}"].font = font(italic=True, size=9, color=MUTED)
    rng = f"C7:I{last - 1}"
    red, gray, green, amber = "FCE1E0", "EDEFF4", "E0F3EA", "C88912"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['C7="OFF"'], fill=PatternFill("solid", fgColor=red), font=Font(color="B42318", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['C7="Feriado"'], fill=PatternFill("solid", fgColor=amber), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['OR(C7="Ausente",C7="Ausente justificado")'], fill=PatternFill("solid", fgColor=gray), font=Font(color="5B6678", bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['C7="Vacaciones"'], fill=PatternFill("solid", fgColor=green), font=Font(color="166534", bold=True)))
    ws.conditional_formatting.add("C4:I6", FormulaRule(formula=['C$6="Feriado"'], fill=PatternFill("solid", fgColor=amber), font=Font(color="FFFFFF", bold=True)))


def team_sheet(ws, team, example):
    ws.title = "Equipo"
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 5, "B": 22, "C": 16, "D": 12, "E": 18, "F": 34}.items():
        ws.column_dimensions[c].width = w
    ws.merge_cells("A1:F1")
    ws["A1"] = "Equipo"
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Cargá a cada persona una sola vez. La planilla semanal toma los nombres de acá (hasta 10)."
    ws["A2"].font = font(size=10, color=MUTED)
    for col, txt in zip("ABCDEF", ["#", "Colaborador", "Puesto", "Cobra por", "OFF (horas que paga)", "Notas"]):
        cell = ws[f"{col}4"]
        cell.value, cell.fill, cell.font, cell.border = txt, HEAD, font(bold=True, color="FFFFFF", size=10), BOX
    dv = DataValidation(type="list", formula1='"Horas,Turnos"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Cobra por", error="Elegí Horas o Turnos.")
    ws.add_data_validation(dv)
    for i in range(PEOPLE):
        r = 5 + i
        ws[f"A{r}"] = i + 1
        ws[f"A{r}"].font = font(color=MUTED)
        for col in "BCDEF":
            ws[f"{col}{r}"].fill, ws[f"{col}{r}"].border, ws[f"{col}{r}"].font = INPUT, BOX, font()
        dv.add(ws[f"D{r}"])
        if i < len(team):
            p = team[i]
            ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"] = p["nombre"], p["puesto"], p["cobra"], p.get("nota", "")
            if p["cobra"] == "Horas":
                ws[f"E{r}"] = hm(p["off"])
                ws[f"E{r}"].number_format = "[h]:mm"
            else:
                ws[f"E{r}"] = 1
                ws[f"E{r}"].number_format = '0" turno"'
    ws["B16"] = ("«OFF» solo se usa en horas para quienes cobran por horas (ej. 4:30). Quienes cobran por turnos suman "
                 "1 turno por OFF.")
    ws["B16"].font = font(size=9, italic=True, color=MUTED)
    if example:
        ws["B17"] = "La fila 1 es un ejemplo del formato: reemplazala por tu equipo."
        ws["B17"].font = font(size=9, italic=True, color="B42318")


def summary_sheet(ws):
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 20, "B": 13, "C": 10, "D": 13, "E": 16, "F": 12, "G": 11, "H": 12, "I": 16, "J": 26}.items():
        ws.column_dimensions[c].width = w
    ws.merge_cells("A1:J1")
    ws["A1"] = "Resumen de la semana"
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    ws["A2"], ws["B2"], ws["C2"], ws["D2"], ws["E2"] = "Semana", "=Semana!I2", "al", "=Semana!K2", "=Semana!E2"
    for ref in ("B2", "D2"):
        ws[ref].number_format = "dd/mm/yyyy"
    for ref in ("A2", "B2", "C2", "D2", "E2"):
        ws[ref].font = font(bold=True)
    heads = ["Colaborador", "Puesto", "Cobra por", "Trabajado", "OFF", "Feriado", "Ausentes", "Vacaciones", "Total a liquidar", "Observaciones"]
    for col, txt in zip("ABCDEFGHIJ", heads):
        cell = ws[f"{col}4"]
        cell.value, cell.fill, cell.font, cell.border = txt, HEAD, font(bold=True, color="FFFFFF", size=10), BOX
        cell.alignment = Alignment(horizontal="center")
    for i in range(PEOPLE):
        r, s, n = 5 + i, 7 + 7 * i, i + 1
        st = f"Semana!C{s + 4}:I{s + 4}"
        ws[f"A{r}"] = f"=Semana!A{s}"
        ws[f"B{r}"] = f'=IF(A{r}="","",INDEX(Equipo!$C$5:$C$14,{n})&"")'
        ws[f"C{r}"] = f'=IF(A{r}="","",INDEX(Equipo!$D$5:$D$14,{n})&"")'
        ws[f"D{r}"] = (f'=IF(A{r}="","",IF(C{r}="Horas",TEXT(Semana!J{s},"[h]:mm")&" hs",'
                       f'Semana!K{s}&IF(Semana!K{s}=1," turno"," turnos")))')
        ws[f"E{r}"] = (f'=IF(A{r}="","",IF(Semana!L{s}=0,"–",IF(C{r}="Horas",TEXT(Semana!L{s}*Semana!R{s},"[h]:mm")&" hs ("&'
                       f'Semana!L{s}&IF(Semana!L{s}=1," día)"," días)"),Semana!L{s}&IF(Semana!L{s}=1," turno"," turnos"))))')
        ws[f"F{r}"] = f'=IF(A{r}="","",Semana!M{s})'
        ws[f"G{r}"] = f'=IF(A{r}="","",IF(Semana!N{s}=0,"–",Semana!N{s}&IF(Semana!N{s}=1," día"," días")))'
        ws[f"H{r}"] = f'=IF(A{r}="","",IF(COUNTIF({st},"Vacaciones")=0,"–",COUNTIF({st},"Vacaciones")&" días"))'
        ws[f"I{r}"] = f'=IF(A{r}="","",Semana!O{s})'
        for col in "ABCDEFGHIJ":
            cell = ws[f"{col}{r}"]
            cell.border = BOX
            cell.font = font(bold=col in "AI")
            cell.alignment = Alignment(horizontal="left" if col in "AJ" else "center")
        ws[f"I{r}"].fill = TOTAL
        ws[f"J{r}"].fill = INPUT
    ws["A16"] = ("Trabajado, OFF y feriado se suman en «Total a liquidar» (en horas o turnos según cómo cobra cada persona). "
                 "Ausentes y vacaciones se informan en días.")
    ws["A16"].font = font(size=9, italic=True, color=MUTED)
    ws.freeze_panes = "A5"


def holidays_sheet(ws, rows):
    ws.title = "Feriados"
    ws.sheet_view.showGridLines = False
    for c, w in {"A": 13, "B": 60, "C": 14}.items():
        ws.column_dimensions[c].width = w
    ws.merge_cells("A1:C1")
    ws["A1"] = "Feriados"
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    ws["A2"] = ("Los feriados de esta lista se pintan en ámbar en el encabezado de la semana como recordatorio. "
                "Para pagarlo, elegí «Feriado» en el Estado de quien trabajó ese día.")
    ws["A2"].font = font(size=9, color=MUTED)
    for col, txt in zip("ABC", ["Fecha", "Motivo", "Tipo"]):
        cell = ws[f"{col}4"]
        cell.value, cell.fill, cell.font, cell.border = txt, HEAD, font(bold=True, color="FFFFFF", size=10), BOX
    for i, (d, motivo, tipo) in enumerate(rows):
        r = 5 + i
        ws[f"A{r}"] = date.fromisoformat(d)
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        ws[f"B{r}"], ws[f"C{r}"] = motivo, tipo
        for col in "ABC":
            ws[f"{col}{r}"].border, ws[f"{col}{r}"].font = BOX, font()
    ws.freeze_panes = "A5"


def guide_sheet(ws, demo):
    ws.title = "Cómo usar"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 110
    ws.merge_cells("A1:C1")
    ws["A1"] = "Planilla semanal de horarios · GESTLOK"
    ws["A1"].font = font(bold=True, size=16, color=NAVY)
    ws["B2"] = ("EJEMPLO CON DATOS FICTICIOS (demo de GESTLOK). Ninguna persona ni local es real."
                if demo else "Plantilla predeterminada de GESTLOK. También podés usar tu propia planilla.")
    ws["B2"].font = font(bold=True, color="B42318" if demo else NAVY)
    steps = [
        "Pestaña «Equipo»: cargá a tu gente una sola vez (nombre, puesto, si cobra por horas o por turnos y cuántas horas vale un OFF).",
        "Pestaña «Semana»: escribí el local y la fecha del lunes. Los días, las fechas y los feriados se completan solos.",
        "Cargá entrada y salida con formato 10:00. Si trabaja de mañana y de noche, cargá los dos turnos. Si sale pasada la medianoche (ej. 01:00), se calcula bien.",
        "En «Estado» elegí OFF, Feriado, Ausente, Ausente justificado o Vacaciones. «Feriado» es cuando la persona trabajó ese día feriado. Si trabajó un día normal, dejalo vacío.",
        "Pestaña «Resumen»: total de la semana de cada persona (trabajado, OFF, feriado, ausentes y vacaciones), listo para imprimir o enviar.",
        "Nueva semana: duplicá el archivo o la pestaña, borrá los horarios celestes y cambiá la fecha del lunes.",
        "Para usarla en GESTLOK: subila a tu Google Drive, abrila con Google Sheets y pegá su enlace en Administrar a los empleados › Horarios.",
    ]
    for i, txt in enumerate(steps):
        ws[f"B{4 + i}"] = i + 1
        ws[f"C{4 + i}"] = txt
        ws[f"B{4 + i}"].font = font(bold=True, color=NAVY)
        ws[f"C{4 + i}"].font = font()
        ws[f"C{4 + i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B12"] = "Colores"
    ws["B12"].font = font(bold=True, color=NAVY)
    colors = [("Celeste: celdas para completar", INPUT), ("Blanco: se calcula solo", None),
              ("Ámbar: feriado", PatternFill("solid", fgColor="C88912")), ("Rojo: OFF", PatternFill("solid", fgColor="FCE1E0")),
              ("Gris: ausente", PatternFill("solid", fgColor="EDEFF4")), ("Verde: vacaciones", PatternFill("solid", fgColor="E0F3EA"))]
    for i, (txt, fill) in enumerate(colors):
        ws[f"C{13 + i}"] = txt
        ws[f"C{13 + i}"].font = font()
        if fill:
            ws[f"B{13 + i}"].fill = fill
    ws["B20"] = "Cómo se calcula"
    ws["B20"].font = font(bold=True, color=NAVY)
    calc = [
        "Horas del día (quienes cobran por horas) = (salida − entrada) del turno mañana + la del turno noche.",
        "Turnos (quienes cobran por turnos): cada horario de entrada cargado suma 1 turno, de mañana o de noche.",
        "OFF: por horas suma las «horas que paga» de la pestaña Equipo; por turnos suma 1 turno.",
        "Feriado: se paga aparte, además de lo trabajado (horas o turnos de ese día).",
        "Total a liquidar = Trabajado + OFF + Feriado. Ausentes y vacaciones se informan en días.",
    ]
    for i, txt in enumerate(calc):
        ws[f"C{21 + i}"] = txt
        ws[f"C{21 + i}"].font = font()
        ws[f"C{21 + i}"].alignment = Alignment(wrap_text=True)


def build(path, *, demo):
    wb = Workbook()
    guide_sheet(wb.active, demo)
    if demo:
        team, week, local, monday = DEMO["equipo"], DEMO["semana"], DEMO["local"], date.fromisoformat(DEMO["lunes"])
    else:
        team = [{"nombre": "Nombre de ejemplo", "puesto": "Staff", "cobra": "Horas", "off": "4:30", "nota": "Fila de ejemplo"}]
        week = {"Nombre de ejemplo": [["10:00", "15:00", "", "", ""], ["", "", "", "", "OFF"]]}
        local, monday = "Tu local", date(2026, 11, 2)
    week_sheet(wb.create_sheet(), local, monday, team, week)
    summary_sheet(wb.create_sheet())
    team_sheet(wb.create_sheet(), team, example=not demo)
    holidays_sheet(wb.create_sheet(), DEMO["feriados"])
    for ws in wb.worksheets:
        # Para imprimir o pasar a PDF: horizontal y a lo ancho de una hoja.
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.active = 1
    wb.properties.creator = "GESTLOK"
    wb.properties.title = "Planilla semanal de horarios"
    wb.save(path)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build(OUT / "GESTLOK-plantilla-horarios.xlsx", demo=False)
    build(OUT / "GESTLOK-demo-horarios-ficticio.xlsx", demo=True)
    print("ok")
